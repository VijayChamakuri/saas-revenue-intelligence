"""Build excel/finance_reconciliation.xlsx from the certified warehouse marts.

Warehouse values are written as inputs; every check, variance, flag, retention ratio and
scenario output is a live Excel formula. No macros. Formula cells are locked, scenario input
cells are unlocked and shaded, and the Control sheet records a run ID and export hashes.
"""

from __future__ import annotations

import argparse
import hashlib
import subprocess
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

SHEETS = [
    "Control", "MRR Bridge", "Billing Reconciliation", "Exceptions",
    "Scenario Inputs", "Scenario Output", "Metric Definitions",
]
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
MONEY = '#,##0.00;[Red]-#,##0.00'
PERCENT = "0.0%"
CAVEAT = "Synthetic data from a seeded generator. Not real company, customer or financial results."
TITLE = "SaaS Revenue Intelligence | MRR, Retention & Billing Controls"

BRIDGE_SQL = """
select b.month_start, b.opening_mrr, b.new_mrr, b.expansion_mrr, b.reactivation_mrr,
       b.contraction_mrr, b.churned_mrr, b.closing_mrr
from analytics_revenue.mart_mrr_bridge b order by 1
"""
BILLING_SQL = """
select date_trunc('month', invoice_date)::date as month_start, count(*) as invoices,
       sum(total_amount) as invoiced, sum(successful_payments) as paid,
       sum(failed_payment_exposure) as unpaid, sum(refunded_amount) as refunded,
       sum(net_collected_cash) as net_cash, sum(recognized_amount) as recognized,
       sum(deferred_amount) as deferred
from analytics_finance.fct_billing_reconciliation group by 1 order by 1
"""
EXCEPTION_SQL = """
select e.invoice_id, e.exception_type, e.severity, r.invoice_date::date as invoice_date,
       r.customer_id, r.total_amount as billed_amount, e.details
from analytics_finance.mart_finance_exceptions e
left join analytics_finance.fct_billing_reconciliation r using (invoice_id)
order by e.severity, r.invoice_date, e.invoice_id
"""
EXPORTS = ["mrr_bridge", "billing_reconciliation", "finance_exceptions"]

DEFINITIONS = [
    ("MRR", "Sum of active subscription monthly recurring revenue at month end", "Closing MRR", "mart_mrr_bridge"),
    ("Net new MRR", "New + expansion + reactivation - contraction - churned", "C+D+E-F-G", "mart_mrr_bridge"),
    ("NRR", "(Opening + expansion + reactivation - contraction - churned) / opening", "MRR Bridge column M", "mart_revenue_kpis"),
    ("GRR", "(Opening - contraction - churned) / opening", "MRR Bridge column N", "mart_revenue_kpis"),
    ("Tie-out variance", "Invoiced - successful payments - unpaid failed exposure", "Billing Reconciliation column J", "fct_billing_reconciliation"),
    ("Net cash", "Successful payments - refunds", "Billing Reconciliation column K", "fct_billing_reconciliation"),
    ("Failed-payment exposure", "Invoice total on invoices without a successful payment", "Billing Reconciliation column E", "fct_billing_reconciliation"),
    ("Finance exception", "Invoice-level rule violation queued for review", "Exceptions sheet", "mart_finance_exceptions"),
]


@dataclass(frozen=True)
class WorkbookData:
    bridge: pd.DataFrame
    billing: pd.DataFrame
    exceptions: pd.DataFrame
    export_hashes: dict[str, str]


def load_data(database: Path, exports: Path) -> WorkbookData:
    with duckdb.connect(str(database), read_only=True) as con:
        bridge, billing = con.execute(BRIDGE_SQL).df(), con.execute(BILLING_SQL).df()
        exceptions = con.execute(EXCEPTION_SQL).df()
    hashes = {
        name: hashlib.sha256((exports / f"{name}.csv").read_bytes()).hexdigest()
        for name in EXPORTS if (exports / f"{name}.csv").exists()
    }
    return WorkbookData(bridge, billing, exceptions, hashes)


def _header(ws: Worksheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=column, value=label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32


def _table(ws: Worksheet, name: str, first_col: str, header_row: int, last_col: str, last_row: int) -> None:
    """Register a named Excel Table (header row plus data) so the range filters and sorts as one object."""
    table = Table(displayName=name, ref=f"{first_col}{header_row}:{last_col}{max(last_row, header_row + 1)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)


def _print_setup(ws: Worksheet, refresh: str, one_page: bool = False, title_rows: str | None = None) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    if title_rows:
        ws.print_title_rows = title_rows
    ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    header, footer = ws.oddHeader, ws.oddFooter
    assert header is not None and footer is not None
    header.left.text, header.left.size = TITLE.replace("&", "&&"), 9
    header.right.text, header.right.size = ws.title, 9
    footer.left.text, footer.left.size = "Synthetic data. Not real financial results.", 8
    footer.center.text, footer.center.size = refresh, 8
    footer.right.text, footer.right.size = "Page &P of &N", 8


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _bridge_sheet(ws: Worksheet, frame: pd.DataFrame) -> int:
    ws["A1"], ws["A1"].font = "MRR bridge and roll-forward checks", TITLE_FONT
    ws["A2"] = "Columns B to H are warehouse values. Columns I to P are formulas."
    _header(ws, 4, ["Month", "Opening", "New", "Expansion", "Reactivation", "Contraction", "Churned",
                    "Closing (warehouse)", "Net new", "Closing (calculated)", "Variance", "Bridge check",
                    "NRR", "GRR", "ARR", "Opening ties to prior close"])
    first = 5
    for offset, row in enumerate(frame.itertuples(index=False)):
        r = first + offset
        ws.cell(r, 1, row.month_start).number_format = "yyyy-mm"
        for column, value in enumerate(row[1:8], start=2):
            ws.cell(r, column, float(value)).number_format = MONEY
        ws.cell(r, 9, f"=C{r}+D{r}+E{r}-F{r}-G{r}").number_format = MONEY
        ws.cell(r, 10, f"=B{r}+I{r}").number_format = MONEY
        ws.cell(r, 11, f"=ROUND(J{r}-H{r},2)").number_format = MONEY
        ws.cell(r, 12, f'=IF(ABS(K{r})<=0.01,"OK","CHECK")')
        ws.cell(r, 13, f'=IF(B{r}>0,(B{r}+D{r}+E{r}-F{r}-G{r})/B{r},"")').number_format = PERCENT
        ws.cell(r, 14, f'=IF(B{r}>0,(B{r}-F{r}-G{r})/B{r},"")').number_format = PERCENT
        ws.cell(r, 15, f"=H{r}*12").number_format = MONEY
        ws.cell(r, 16, '="OK"' if offset == 0 else f'=IF(ABS(B{r}-H{r - 1})<=0.01,"OK","CHECK")')
    _widths(ws, [10] + [15] * 15)
    ws.freeze_panes = "B5"
    _table(ws, "MrrBridge", "A", 4, "P", first + len(frame) - 1)
    return first + len(frame) - 1


def _billing_sheet(ws: Worksheet, frame: pd.DataFrame) -> int:
    ws["A1"], ws["A1"].font = "Invoice, payment, refund and revenue tie-out", TITLE_FONT
    ws["A2"] = "Invoices must equal successful payments plus unpaid exposure. Recognized plus deferred must equal invoiced."
    _header(ws, 4, ["Invoice month", "Invoices", "Invoiced", "Paid", "Unpaid (failed)", "Refunded",
                    "Net cash (warehouse)", "Recognized", "Deferred", "Tie-out variance",
                    "Net cash (calculated)", "Cash variance", "Revenue variance", "Reconciliation check"])
    first = 5
    for offset, row in enumerate(frame.itertuples(index=False)):
        r = first + offset
        ws.cell(r, 1, row.month_start).number_format = "yyyy-mm"
        ws.cell(r, 2, int(row.invoices))
        # warehouse order: invoiced, paid, unpaid, refunded, net_cash, recognized, deferred
        for column, value in zip((3, 4, 5, 6, 7, 8, 9), row[2:9], strict=True):
            ws.cell(r, column, float(value)).number_format = MONEY
        ws.cell(r, 10, f"=ROUND(C{r}-D{r}-E{r},2)").number_format = MONEY
        ws.cell(r, 11, f"=D{r}-F{r}").number_format = MONEY
        ws.cell(r, 12, f"=ROUND(K{r}-G{r},2)").number_format = MONEY
        ws.cell(r, 13, f"=ROUND(C{r}-H{r}-I{r},2)").number_format = MONEY
        ws.cell(r, 14, f'=IF(AND(ABS(J{r})<=0.01,ABS(L{r})<=0.01,ABS(M{r})<=0.01),"OK","CHECK")')
    last = first + len(frame) - 1
    _table(ws, "BillingReconciliation", "A", 4, "N", last)
    total = last + 1
    ws.cell(total, 1, "Total").font = Font(bold=True)
    for column in range(2, 14):
        letter = get_column_letter(column)
        cell = ws.cell(total, column, f"=SUM({letter}{first}:{letter}{last})")
        cell.font = Font(bold=True)
        cell.number_format = "#,##0" if column == 2 else MONEY
    _widths(ws, [13] + [16] * 13)
    ws.freeze_panes = "B5"
    return last


def _exception_sheet(ws: Worksheet, frame: pd.DataFrame) -> None:
    ws["A1"], ws["A1"].font = "Finance exception queue", TITLE_FONT
    ws["A2"] = "Summary formulas read the detail list below. Billed amount counts an invoice once per rule it breaks."
    rules = sorted(frame.exception_type.unique())
    _header(ws, 4, ["Rule", "Severity", "Records", "Billed amount"])
    detail_header = 6 + len(rules) + 1
    first, last = detail_header + 1, detail_header + len(frame)
    for offset, rule in enumerate(rules):
        r = 5 + offset
        severity = frame.loc[frame.exception_type == rule, "severity"].iloc[0]
        ws.cell(r, 1, rule)
        ws.cell(r, 2, severity)
        ws.cell(r, 3, f"=COUNTIF($B${first}:$B${last},A{r})")
        ws.cell(r, 4, f"=SUMIF($B${first}:$B${last},A{r},$F${first}:$F${last})").number_format = MONEY
    total = 5 + len(rules)
    _table(ws, "ExceptionSummary", "A", 4, "D", total - 1)
    ws.cell(total, 1, "Total").font = Font(bold=True)
    ws.cell(total, 3, f"=SUM(C5:C{total - 1})").font = Font(bold=True)
    ws.cell(total, 4, f"=SUM(D5:D{total - 1})").number_format = MONEY
    _header(ws, detail_header, ["Invoice", "Rule", "Severity", "Invoice date", "Customer", "Billed amount", "Details"])
    for offset, row in enumerate(frame.itertuples(index=False)):
        r = first + offset
        ws.cell(r, 1, row.invoice_id)
        ws.cell(r, 2, row.exception_type)
        ws.cell(r, 3, row.severity)
        if pd.notna(row.invoice_date):
            ws.cell(r, 4, row.invoice_date).number_format = "yyyy-mm-dd"
        ws.cell(r, 5, row.customer_id)
        if pd.notna(row.billed_amount):
            ws.cell(r, 6, float(row.billed_amount)).number_format = MONEY
        ws.cell(r, 7, row.details)
    _widths(ws, [30, 26, 12, 16, 14, 16, 60])
    ws.freeze_panes = f"A{first}"
    _table(ws, "ExceptionDetail", "A", detail_header, "G", last)


def _scenario_inputs(ws: Worksheet) -> dict[str, str]:
    ws["A1"], ws["A1"].font = "Scenario inputs", TITLE_FONT
    ws["A2"] = "Shaded cells are the only editable cells in this workbook. Outputs are simulated, not observed."
    _header(ws, 4, ["Assumption", "Value", "Meaning"])
    rows = [
        ("Churn reduction", 0.10, "Share of churned MRR avoided each month"),
        ("Pricing change", 0.03, "Uplift applied to closing MRR"),
        ("Expansion uplift", 0.10, "Extra expansion as a share of observed expansion MRR"),
        ("Failed-payment recovery", 0.20, "Share of unpaid failed exposure recovered as cash"),
    ]
    cells = {}
    for offset, (label, value, meaning) in enumerate(rows):
        r = 5 + offset
        ws.cell(r, 1, label)
        cell = ws.cell(r, 2, value)
        cell.fill, cell.number_format = INPUT_FILL, PERCENT
        cell.protection = Protection(locked=False)
        ws.cell(r, 3, meaning)
        cells[label] = f"'Scenario Inputs'!$B${r}"
    _widths(ws, [28, 12, 60])
    _table(ws, "ScenarioInputs", "A", 4, "C", 4 + len(rows))
    return cells


def _scenario_output(ws: Worksheet, inputs: dict[str, str], bridge_last: int, billing_last: int) -> None:
    ws["A1"], ws["A1"].font = "Scenario output (simulated)", TITLE_FONT
    ws["A2"] = "Formulas apply the shaded assumptions to observed monthly values. Not a forecast."
    _header(ws, 4, ["Month", "Baseline MRR", "Avoided churn", "Extra expansion", "Pricing effect",
                    "Scenario MRR", "Incremental MRR", "Baseline net cash", "Recovered cash", "Scenario net cash"])
    for offset in range(bridge_last - 4):
        r, b = 5 + offset, 5 + offset
        ws.cell(r, 1, f"='MRR Bridge'!A{b}").number_format = "yyyy-mm"
        ws.cell(r, 2, f"='MRR Bridge'!H{b}")
        ws.cell(r, 3, f"='MRR Bridge'!G{b}*{inputs['Churn reduction']}")
        ws.cell(r, 4, f"='MRR Bridge'!D{b}*{inputs['Expansion uplift']}")
        ws.cell(r, 5, f"='MRR Bridge'!H{b}*{inputs['Pricing change']}")
        ws.cell(r, 6, f"=B{r}+C{r}+D{r}+E{r}")
        ws.cell(r, 7, f"=F{r}-B{r}")
        if b <= billing_last:
            ws.cell(r, 8, f"='Billing Reconciliation'!G{b}")
            ws.cell(r, 9, f"='Billing Reconciliation'!E{b}*{inputs['Failed-payment recovery']}")
            ws.cell(r, 10, f"=H{r}+I{r}")
        for column in (2, 3, 4, 5, 6, 7, 8, 9, 10):
            ws.cell(r, column).number_format = MONEY
    _widths(ws, [10] + [17] * 9)
    ws.freeze_panes = "B5"
    _table(ws, "ScenarioOutput", "A", 4, "J", 4 + (bridge_last - 4))


def _source_commit() -> str:
    try:
        return subprocess.run(["git", "rev-parse", "--short=12", "HEAD"], capture_output=True, text=True,
                              check=True).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def _control(ws: Worksheet, data: WorkbookData, bridge_last: int, billing_last: int, refreshed: str) -> None:
    run_id = hashlib.sha256("".join(f"{k}{v}" for k, v in sorted(data.export_hashes.items())).encode()).hexdigest()[:12]
    detail_header = 6 + data.exceptions["exception_type"].nunique() + 1  # same layout as _exception_sheet
    detail_first, detail_last = detail_header + 1, detail_header + len(data.exceptions)
    ws["A1"], ws["A1"].font = TITLE, TITLE_FONT
    ws["A2"], ws["A2"].font = CAVEAT, Font(bold=True, color="9C0006")
    ws["A3"] = "Built by src/excel/build_workbook.py from the tested dbt marts. No macros. Shaded cells on Scenario Inputs are the only editable cells."
    rows = [
        ("Refreshed (UTC)", refreshed),
        ("Source commit", _source_commit()),
        ("Run ID (hash of exports)", run_id),
        ("Data window", f"{data.bridge.month_start.min():%Y-%m} to {data.bridge.month_start.max():%Y-%m}"),
        ("MRR bridge breaks", f"=COUNTIF('MRR Bridge'!L5:L{bridge_last},\"CHECK\")"),
        ("Opening balance breaks", f"=COUNTIF('MRR Bridge'!P5:P{bridge_last},\"CHECK\")"),
        ("Billing reconciliation breaks", f"=COUNTIF('Billing Reconciliation'!N5:N{billing_last},\"CHECK\")"),
        ("Reconciliation status", '=IF(B9+B10+B11=0,"ALL CHECKS PASS","REVIEW")'),
        ("Latest closing MRR", f"='MRR Bridge'!H{bridge_last}"),
        ("Latest net revenue retention", f"='MRR Bridge'!M{bridge_last}"),
        ("Latest gross revenue retention", f"='MRR Bridge'!N{bridge_last}"),
        ("Total invoiced", f"='Billing Reconciliation'!C{billing_last + 1}"),
        ("Total failed-payment exposure", f"='Billing Reconciliation'!E{billing_last + 1}"),
        ("Finance exceptions queued", f"=COUNTA(Exceptions!A{detail_first}:A{detail_last})"),
        ("Caveat", "December 2025 churn is concentrated on the last day of the data window; treat it as a data-dating question."),
    ]
    _header(ws, 4, ["Control item", "Value"])
    for offset, (label, value) in enumerate(rows):
        r = 5 + offset
        ws.cell(r, 1, label)
        cell = ws.cell(r, 2, value)
        if label.startswith(("Latest closing", "Total")):
            cell.number_format = MONEY
        elif "retention" in label:
            cell.number_format = PERCENT
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    ws["B12"].font = Font(bold=True)
    last = 4 + len(rows)
    _table(ws, "ControlItems", "A", 4, "B", last)
    hashes = last + 2
    _header(ws, hashes, ["Export", "SHA-256"])
    for offset, (name, digest) in enumerate(sorted(data.export_hashes.items())):
        ws.cell(hashes + 1 + offset, 1, f"{name}.csv")
        ws.cell(hashes + 1 + offset, 2, digest)
    _table(ws, "ExportHashes", "A", hashes, "B", hashes + len(data.export_hashes))
    _widths(ws, [34, 90])

def _definitions(ws: Worksheet) -> None:
    ws["A1"], ws["A1"].font = "Metric definitions (implemented metrics only)", TITLE_FONT
    _header(ws, 3, ["Metric", "Definition", "Where in this workbook", "Source mart"])
    for offset, row in enumerate(DEFINITIONS):
        for column, value in enumerate(row, start=1):
            ws.cell(4 + offset, column, value)
    _widths(ws, [24, 70, 34, 30])
    _table(ws, "MetricDefinitions", "A", 3, "D", 3 + len(DEFINITIONS))


def build_workbook(data: WorkbookData, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]  # openpyxl always creates one default sheet
    sheets = {name: wb.create_sheet(name) for name in SHEETS}
    bridge_last = _bridge_sheet(sheets["MRR Bridge"], data.bridge)
    billing_last = _billing_sheet(sheets["Billing Reconciliation"], data.billing)
    _exception_sheet(sheets["Exceptions"], data.exceptions)
    inputs = _scenario_inputs(sheets["Scenario Inputs"])
    _scenario_output(sheets["Scenario Output"], inputs, bridge_last, billing_last)
    refreshed = datetime.now(UTC).strftime("%Y-%m-%d %H:%M")
    _control(sheets["Control"], data, bridge_last, billing_last, refreshed)
    _definitions(sheets["Metric Definitions"])
    for ws in wb.worksheets:
        ws.protection.sheet = True  # formulas locked; only shaded scenario inputs are editable
        ws.protection.autoFilter = False  # filters and sorting stay usable on the protected sheets
        ws.protection.sort = False
        header_rows = {"MRR Bridge": "4:4", "Billing Reconciliation": "4:4", "Scenario Output": "4:4"}
        _print_setup(ws, f"Refreshed {refreshed} UTC", one_page=ws.title == "Control", title_rows=header_rows.get(ws.title))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_control_pdf(workbook: Path, output: Path) -> Path | None:
    """Render the workbook with LibreOffice and keep page 1 (the Control sheet) as a one-page PDF."""
    import shutil
    import tempfile

    soffice = shutil.which("soffice") or ("/Applications/LibreOffice.app/Contents/MacOS/soffice"
                                          if Path("/Applications/LibreOffice.app").exists() else None)
    if soffice is None or shutil.which("pdfseparate") is None:
        print("LibreOffice or pdfseparate not found: control PDF not exported")
        return None
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--convert-to", "pdf",
                        "--outdir", tmp, str(workbook)], check=True, capture_output=True, timeout=180)
        output.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run(["pdfseparate", "-f", "1", "-l", "1", str(Path(tmp) / f"{workbook.stem}.pdf"), str(output)],
                       check=True)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database", type=Path, default=Path("data/warehouse/subscription.duckdb"))
    parser.add_argument("--exports", type=Path, default=Path("data/exports"))
    parser.add_argument("--output", type=Path, default=Path("excel/finance_reconciliation.xlsx"))
    parser.add_argument("--pdf", type=Path, default=Path("reports/finance_control_summary.pdf"))
    args = parser.parse_args()
    path = build_workbook(load_data(args.database, args.exports), args.output)
    print(path)
    print(export_control_pdf(path, args.pdf))
