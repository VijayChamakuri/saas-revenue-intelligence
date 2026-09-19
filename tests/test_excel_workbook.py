"""Structural checks on the finance workbook: sheets, formulas, references, protection."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.excel.build_workbook import SHEETS, WorkbookData, build_workbook


def _data() -> WorkbookData:
    months = pd.to_datetime(["2024-01-01", "2024-02-01"])
    bridge = pd.DataFrame({
        "month_start": months, "opening_mrr": [0.0, 100.0], "new_mrr": [100.0, 20.0],
        "expansion_mrr": [0.0, 10.0], "reactivation_mrr": [0.0, 0.0],
        "contraction_mrr": [0.0, 5.0], "churned_mrr": [0.0, 15.0], "closing_mrr": [100.0, 110.0],
    })
    billing = pd.DataFrame({
        "month_start": months, "invoices": [2, 2], "invoiced": [100.0, 110.0], "paid": [100.0, 100.0],
        "unpaid": [0.0, 10.0], "refunded": [0.0, 5.0], "net_cash": [100.0, 95.0],
        "recognized": [100.0, 110.0], "deferred": [0.0, 0.0],
    })
    exceptions = pd.DataFrame({
        "invoice_id": ["I1", "I2"], "exception_type": ["contract_date_conflict"] * 2,
        "severity": ["medium"] * 2, "invoice_date": months, "customer_id": ["C1", "C2"],
        "billed_amount": [40.0, 60.0], "details": ["d1", "d2"],
    })
    return WorkbookData(bridge, billing, exceptions, {"mrr_bridge": "ab" * 32})


@pytest.fixture()
def workbook(tmp_path: Path):
    return load_workbook(build_workbook(_data(), tmp_path / "wb.xlsx"))


def test_expected_sheets_in_order(workbook) -> None:
    assert workbook.sheetnames == SHEETS


def test_formulas_reference_existing_sheets_and_no_macros(workbook, tmp_path: Path) -> None:
    formulas = [c.value for ws in workbook.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 50
    referenced = {m for f in formulas for m in re.findall(r"'([^']+)'!", f)}
    assert referenced and referenced <= set(workbook.sheetnames)
    with zipfile.ZipFile(tmp_path / "wb.xlsx") as package:
        assert not any("vbaProject" in name for name in package.namelist())


def test_bridge_formulas_and_tie_out_checks_exist(workbook) -> None:
    bridge = workbook["MRR Bridge"]
    assert bridge["I5"].value == "=C5+D5+E5-F5-G5"
    assert bridge["K5"].value == "=ROUND(J5-H5,2)"
    assert bridge["P6"].value.startswith("=IF(ABS(B6-H5)")
    billing = workbook["Billing Reconciliation"]
    assert billing["J5"].value == "=ROUND(C5-D5-E5,2)"
    assert billing["M5"].value == "=ROUND(C5-H5-I5,2)"
    assert "ALL CHECKS PASS" in workbook["Control"]["B8"].value


def test_only_scenario_inputs_are_editable(workbook) -> None:
    unlocked = [f"{ws.title}!{c.coordinate}" for ws in workbook.worksheets for row in ws.iter_rows()
                for c in row if c.value is not None and not c.protection.locked]
    assert sorted(unlocked) == [f"Scenario Inputs!B{r}" for r in range(5, 9)]
    assert all(ws.protection.sheet for ws in workbook.worksheets)


def test_scenario_output_reads_inputs(workbook) -> None:
    assert "'Scenario Inputs'!$B$5" in workbook["Scenario Output"]["C5"].value
    assert workbook["Control"]["B4"].value  # run id recorded
